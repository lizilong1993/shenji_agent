#!/usr/bin/env python3
"""Extract the authoritative 庙算 rule tables into JSON (P0).

Inputs (copied from the SDK distribution):
  data/裁决表格.xlsx            consolidated public adjudication workbook
  data/weapons_0.xlsx          direct/indirect weapon range lists
  data/zm_attack_level.xlsx    direct-fire attack level tables to_car/to_peo/to_plane
  data/WpData.csv              master weapon x TarType x ObjBlood x distance table
  data/ele_rect_to_car|peo     elevation modifier matrices (8x12)
  data/minefield_*.csv         minefield damage tables

Output: <out>/rules/data/*.json plus manifest.json.
Run:  python3 -I extract_rules.py <data_dir> <out_dir>
"""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def cell_int(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).strip()
    try:
        return int(s)
    except ValueError:
        return s


def grid(ws, max_row=None, max_col=None):
    rows = []
    for r in range(1, (max_row or ws.max_row) + 1):
        row = []
        for c in range(1, (max_col or ws.max_column) + 1):
            row.append(cell_int(ws.cell(r, c).value))
        rows.append(row)
    return rows


def parse_personnel_sheet(ws):
    out = []
    name = None
    att_max = None
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value not in (None, ""):
            name = str(ws.cell(r, 1).value).strip()
        if ws.cell(r, 2).value not in (None, ""):
            att_max = cell_int(ws.cell(r, 2).value)
        blood = cell_int(ws.cell(r, 3).value)
        if name is None or blood is None:
            continue
        values = {}
        for c in range(4, ws.max_column + 1):
            header = cell_int(ws.cell(2, c).value)
            v = cell_int(ws.cell(r, c).value)
            if header is None or v is None or v == "":
                continue
            values[str(header)] = v
        out.append({"name": name, "att_max": att_max, "blood": blood, "values": values})
    return out


def parse_vehicle_sheet(ws):
    out = []
    for r in range(3, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if name in (None, ""):
            continue
        name = str(name).strip()
        att_max = cell_int(ws.cell(r, 2).value)
        values = {}
        for c in range(3, ws.max_column + 1):
            header = cell_int(ws.cell(2, c).value)
            v = cell_int(ws.cell(r, c).value)
            if header is None or v is None or v == "":
                continue
            values[str(header)] = v
        out.append({"name": name, "att_max": att_max, "values": values})
    return out


def parse_zm_sheet(name):
    path = DATA / "zm_attack_level.xlsx"
    wb = load_workbook(path, data_only=True)
    ws = wb[name]
    out = []
    for r in range(2, ws.max_row + 1):
        wid = cell_int(ws.cell(r, 1).value)
        if wid is None or wid == "":
            continue
        rec = {
            "weapon_id": wid,
            "name": str(ws.cell(r, 2).value or "").strip(),
            "att_min": cell_int(ws.cell(r, 3).value),
            "att_max": cell_int(ws.cell(r, 4).value),
        }
        start = 5
        if name == "to_peo":
            rec["blood"] = cell_int(ws.cell(r, 5).value)
            start = 6
        values = {}
        for c in range(start, ws.max_column + 1):
            header = cell_int(ws.cell(1, c).value)
            v = cell_int(ws.cell(r, c).value)
            if header is None or v is None or v == "":
                continue
            values[str(header)] = v
        rec["values"] = values
        out.append(rec)
    return out


def parse_weapons0():
    wb = load_workbook(DATA / "weapons_0.xlsx", data_only=True)
    out = {}
    for sheet in ("zm_weapons", "jm_weapons"):
        ws = wb[sheet]
        rows = []
        for r in range(2, ws.max_row + 1):
            name = ws.cell(r, 1).value
            if name in (None, ""):
                continue
            rows.append({
                "name": str(name).strip(),
                "weapon_id": cell_int(ws.cell(r, 2).value),
                "tar_type": cell_int(ws.cell(r, 3).value),
                "att_min": cell_int(ws.cell(r, 4).value),
                "att_max": cell_int(ws.cell(r, 5).value),
            })
        out[sheet] = rows
    return out


def parse_wpdata():
    path = DATA / "WpData.csv"
    rows = []
    with open(path, encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        header = next(reader)
        dist_cols = [int(h) for h in header[11:] if h.strip().isdigit()]
        for raw in reader:
            if len(raw) < 11:
                continue
            name = raw[2].strip()
            if not name:
                continue
            values = {}
            for i, d in enumerate(dist_cols):
                v = raw[11 + i].strip() if 11 + i < len(raw) else ""
                if v:
                    values[str(d)] = int(v)
            rows.append({
                "row_id": raw[0].strip(),
                "combo_id": cell_int(raw[1]),
                "name": name,
                "weapon_id": raw[3].strip(),
                "tar_type": cell_int(raw[4]),
                "obj_blood": cell_int(raw[5]),
                "att_min": cell_int(raw[6]),
                "att_max": cell_int(raw[7]),
                "f0": cell_int(raw[8]),
                "f1": cell_int(raw[9]),
                "f2": cell_int(raw[10]),
                "values": values,
            })
    return rows


def parse_ele_rect(path):
    rows = []
    for line in path.read_text().strip().splitlines():
        vals = [int(x) for x in line.split()]
        if vals:
            rows.append(vals)
    return rows


def parse_minefield(path):
    out = []
    with open(path, encoding="utf-8-sig", newline="") as fh:
        for raw in csv.reader(fh):
            if not raw or not raw[0].strip():
                continue
            out.append([c.strip() for c in raw])
    return out


def write_json(name, obj):
    p = OUT / name
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print("wrote", p, p.stat().st_size, "bytes")


def main():
    global DATA, OUT
    DATA = Path(sys.argv[1])
    OUT = Path(sys.argv[2]) / "rules" / "data"
    OUT.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(DATA / "裁决表格.xlsx", data_only=True)

    write_json("observation.json", {
        "schema": "msgeo.observation.v1",
        "source": "docs/docs.md 观察规则表",
        "table": {
            "infantry": {"infantry": 10, "vehicle": 25, "helicopter": 25, "uav": 1},
            "vehicle": {"infantry": 10, "vehicle": 25, "helicopter": 25, "uav": 1},
            "helicopter": {"infantry": 10, "vehicle": 25, "helicopter": 25, "uav": 1},
            "uav": {"infantry": 2, "vehicle": 2, "helicopter": 0, "uav": 0},
        },
        "notes": [
            "掩蔽目标被观察距离减半；车辆目标在观察者更高高程时不受掩蔽减半",
            "目标位于通视的居民地/丛林格时被观察距离减半",
            "无人机/巡飞弹观察距离为1格表示当前/相邻格",
        ],
    })

    write_json("wpdata.json", {
        "schema": "msgeo.wpdata.v1",
        "source": "Data/weapons/WpData.csv",
        "columns": ["row_id", "combo_id", "name", "weapon_id", "tar_type",
                    "obj_blood", "att_min", "att_max", "f0", "f1", "f2", "values"],
        "rows": parse_wpdata(),
    })

    write_json("zm_attack_level.json", {
        "schema": "msgeo.zm_attack_level.v1",
        "source": "Data/weapons/zm_attack_level.xlsx",
        "to_car": parse_zm_sheet("to_car"),
        "to_peo": parse_zm_sheet("to_peo"),
        "to_plane": parse_zm_sheet("to_plane"),
    })

    write_json("weapons_0.json", {
        "schema": "msgeo.weapons0.v1",
        "source": "Data/weapons/weapons_0.xlsx",
        **parse_weapons0(),
    })

    elevation_xlsx = grid(wb["攻击等级高度差修正"])
    write_json("elevation.json", {
        "schema": "msgeo.elevation.v1",
        "source": "Data/weapons/ele_rect_to_car|peo + 裁决表格.xlsx 攻击等级高度差修正",
        "row_labels": list(range(1, 9)),
        "col_labels": list(range(1, 13)),
        "car": parse_ele_rect(DATA / "ele_rect_to_car"),
        "peo": parse_ele_rect(DATA / "ele_rect_to_peo"),
        "workbook_sheet": elevation_xlsx,
        "row_col_semantics": "UNCONFIRMED",
        "notes": [
            "数据字典推测行=射击方相对高程、列=目标方相对高程；裁决表格表头写的是'高度差/距离'，两种解释矛盾，P0 黑盒实测被 MAC 鉴权阻断",
            "workbook 第7行第12列值为 +1，而 ele_rect_to_car 文件同位置为 -1，疑为表格笔误；以数据文件为准",
        ],
    })

    raw_sheets = {}
    for sheet in ("对车辆单位战斗结果", "车辆战损结果修正",
                  "直瞄武器对人员and步兵轻武器对车辆战斗结果表",
                  "人员战损结果修正", "战斗结果修正", "空战斗结果表",
                  "无校射火力散布裁决", "格内校射火力散布裁决"):
        raw_sheets[sheet] = grid(wb[sheet])
    write_json("crt.json", {
        "schema": "msgeo.crt.v1",
        "source": "裁决表格.xlsx",
        "sheets": raw_sheets,
        "notes": [
            "对车辆单位战斗结果：前5行把(车/班数, 攻击等级)映射到列号，后11行是随机数2..12→结果",
            "车辆战损结果修正：目标/射击状态修正 + 按装甲类型的修正结果表",
        ],
    })

    write_json("attack_levels.json", {
        "schema": "msgeo.attack_levels.v1",
        "source": "裁决表格.xlsx",
        "vs_personnel": parse_personnel_sheet(wb["对人员攻击等级表"]),
        "vs_vehicle": parse_vehicle_sheet(wb["对车辆攻击等级表"]),
        "small_arms_vs_vehicle": parse_personnel_sheet(wb["步兵轻武器对车辆攻击等级表"]),
        "vs_air": {"source": "裁决表格.xlsx 武器攻击等级表",
                   "raw": grid(wb["武器攻击等级表"])},
    })

    write_json("minefield.json", {
        "schema": "msgeo.minefield.v1",
        "source": "Data/weapons/minefield_*.csv",
        "personnel": parse_minefield(DATA / "minefield_damage_to_personnel.csv"),
        "vehicle": parse_minefield(DATA / "minefield_damage_to_vehicle.csv"),
        "vehicle_calibration": parse_minefield(DATA / "minefield_damage_calibration_to_vehicle.csv"),
    })

    write_json("manifest.json", {
        "schema": "msgeo.rules.v1",
        "source_date": "2026-09-15 (SDK data snapshot)",
        "generated": "2026-10-21",
        "files": sorted(p.name for p in OUT.glob("*.json")),
    })


if __name__ == "__main__":
    main()
